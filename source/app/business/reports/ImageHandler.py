#  IRIS Source Code
#  Copyright (C) 2021 - Airbus CyberSecurity (SAS)
#  contact@dfir-iris.org
#  Created by Lukas Zurschmiede @LukyLuke
#
#  This program is free software; you can redistribute it and/or
#  modify it under the terms of the GNU Lesser General Public
#  License as published by the Free Software Foundation; either
#  version 3 of the License, or (at your option) any later version.
#
#  This program is distributed in the hope that it will be useful,
#  but WITHOUT ANY WARRANTY; without even the implied warranty of
#  MERCHANTABILITY or FITNESS FOR A PARTICULAR PURPOSE.  See the GNU
#  Lesser General Public License for more details.
#
#  You should have received a copy of the GNU Lesser General Public License
#  along with this program; if not, write to the Free Software Foundation,
#  Inc., 51 Franklin Street, Fifth Floor, Boston, MA  02110-1301, USA.

import logging
import os
import shutil
import uuid
import re

from pathlib import Path
from urllib.parse import urlsplit, parse_qs
from docxtpl import DocxTemplate

from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx_generator.globals.picture_globals import PictureGlobals
from docx_generator.exceptions.rendering_error import RenderingError

from app.datamgmt.datastore.datastore_db import datastore_get_local_file_path


class ImageHandler(PictureGlobals):
    def __init__(self, template: DocxTemplate, base_path: str):
        self._logger = logging.getLogger(__name__)
        # Width (% of page text-width) requested for the image currently being added, parsed
        # from the &iriswidth= query param. None => embed at native size (capped to page width).
        self._target_width_pct = None
        PictureGlobals.__init__(self, template, base_path)

    def _process_remote(self, image_path: str) -> str:
        """
        Checks if the given Link is a datastore-link and if so, save the image locally for further processing.
        :
        A Datastore Links looks like this: https://localhost:4433/datastore/file/view/2?cid=1
        """
        res = re.search(r'datastore\/file\/view\/(\d+)\?cid=(\d+)', image_path)
        if not res:
            return super()._process_remote(image_path)

        if image_path[:4] == 'http' and len(res.groups()) == 2:
            file_id = res.groups(0)[0]
            case_id = res.groups(0)[1]
            has_error, dsf = datastore_get_local_file_path(file_id, case_id)

            if has_error:
                raise RenderingError(self._logger, f'File-ID {file_id} does not exist in Case {case_id}')
            if not Path(dsf.file_local_name).is_file():
                raise RenderingError(self._logger, f'File {dsf.file_local_name} does not exists on the server. Update or delete virtual entry')

            file_ext = os.path.splitext(dsf.file_original_name)[1]
            file_name = os.path.join(self._output_path, str(uuid.uuid4())) + file_ext
            return_value = shutil.copy(dsf.file_local_name, file_name)
            return return_value
        return super()._process_remote(image_path)

    def add_picture(self, image_path: str, position: str = 'CENTER'):
        """
        Adds a picture, honoring an optional width (percent of page text-width) carried in the
        image URL as ?...&iriswidth=<pct>. The width is read here (render_image only gives us the
        src) and applied in _process_image. Reset in finally so it never leaks to the next image.
        """
        self._target_width_pct = None
        try:
            query = parse_qs(urlsplit(image_path).query)
            if 'iriswidth' in query:
                self._target_width_pct = max(1, min(100, int(query['iriswidth'][0])))
        except (ValueError, TypeError, AttributeError):
            self._target_width_pct = None

        try:
            return super().add_picture(image_path, position)
        except RenderingError as e:
            # Never let a single bad image (missing file, invalid path, decode error, etc.) abort
            # the whole report — skip it and continue, otherwise the route returns a JSON error
            # that the browser saves as a corrupt .docx. Other exceptions (e.g. a filesystem
            # error creating the temp image directory) are real bugs and should surface.
            self._logger.error('Skipping image in report (%s): %s', image_path, e)
            return self._template.new_subdoc()
        finally:
            self._target_width_pct = None

    def _process_image(self, position, image_filename):
        """
        Same as PictureGlobals._process_image, but when a target width % was requested, scale the
        picture to that fraction of the page text-width instead of only scaling down oversized images.
        """
        sub_document = self._template.new_subdoc()

        last_section = sub_document.sections[-1]
        page_width = last_section.page_width - last_section.left_margin - last_section.right_margin

        try:
            picture = sub_document.add_picture(image_filename)
        except Exception as e:
            self._logger.debug('Error while adding image {}: {}'.format(image_filename, e.__str__()))
            raise RenderingError(self._logger, 'Image could not be added (try PNG instead of JPEG): {}'.format(image_filename))

        target_pct = getattr(self, '_target_width_pct', None)
        if target_pct:
            # Percent of the MAIN document's text width (where the image actually renders);
            # the subdoc inherits a default section whose width differs. Fall back to the
            # subdoc width if the main document is unavailable.
            try:
                main_section = self._template.get_docx().sections[-1]
                ref_width = main_section.page_width - main_section.left_margin - main_section.right_margin
            except Exception:
                ref_width = page_width
            self._scale_picture(picture, int(ref_width * target_pct / 100))
        elif picture.width > page_width:
            self._scale_picture(picture, page_width)

        if position in self._available_alignment_values:
            last_paragraph = sub_document.paragraphs[-1]
            last_paragraph.alignment = getattr(WD_PARAGRAPH_ALIGNMENT, position)

        return sub_document
